"""Extrae los catálogos (listas desplegables) de la hoja oculta 'Validador- Datos'
del formato 13-IF-016 y los escribe como src/js/catalogos.js.

Uso:  py -X utf8 tools/extract_catalogos.py
Ejecutar de nuevo cada vez que la SED actualice el Excel original.
"""
import json
import pathlib
import warnings

import openpyxl

warnings.filterwarnings("ignore")

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = next(ROOT.glob("13-IF-016*.xlsx"))
OUT = ROOT / "src" / "js" / "catalogos.js"


def s(v):
    return None if v is None else str(v)


def main():
    wb = openpyxl.load_workbook(XLSX)
    v = wb["Validador- Datos"]
    ded = wb["FORMATO DEDUCCIONES"]

    def col(c1, r1, r2, cols):
        out = []
        for r in range(r1, r2 + 1):
            row = [v.cell(r, c).value for c in cols]
            if all(x is None for x in row):
                continue
            out.append(row)
        return out

    from openpyxl.utils import column_index_from_string as ci

    # Tipo de documento: el nombre se conserva EXACTO (incluye espacios finales)
    # porque las fórmulas VLOOKUP del Excel dependen del texto literal.
    tipos_doc = [
        {"nombre": n, "dian": d, "apot": a}
        for n, d, a in col(None, 2, 14, [ci("K"), ci("L"), ci("M")])
    ]

    paises = [
        {"nombre": n, "codigo": s(c)}
        for n, c in col(None, 2, 249, [ci("T"), ci("U")])
    ]

    departamentos = []
    seen = set()
    for n, c in col(None, 2, 36, [ci("W"), ci("X")]):
        if n in seen:
            continue
        seen.add(n)
        departamentos.append({"nombre": n, "codigo": s(c)})

    municipios = [
        {"depto": d, "codDepto": s(cd), "nombre": n, "codigo": s(cm)}
        for d, cd, n, cm in col(None, 2, 1120, [ci("AB"), ci("AC"), ci("AF"), ci("AG")])
    ]

    bancos = []
    for social, abrev, _nit, cod in col(None, 2, 77, [ci("AN"), ci("AO"), ci("AP"), ci("AQ")]):
        nombre_up = (abrev or "").upper()
        interno = any(k in nombre_up for k in ("DUMMY", "DUMY", "TRANSITORIAS"))
        bancos.append({"nombre": abrev, "social": social, "codigo": s(cod), "interno": interno})

    tipos_cuenta = [{"nombre": n, "codigo": c} for n, c in col(None, 80, 82, [ci("AN"), ci("AO")])]
    regimen = [{"nombre": n, "sigla": c} for n, c in col(None, 22, 23, [ci("K"), ci("L")])]
    areas = [{"nombre": n, "codigo": c} for n, c in col(None, 36, 78, [ci("K"), ci("L")])]

    rit = [
        {"codigo": c, "agrupacion": a, "descripcion": (d or "").strip(), "tarifa": t}
        for c, a, d, t in col(None, 2, 553, [ci("AS"), ci("AT"), ci("AU"), ci("AV")])
    ]
    ciiu = [
        {"codigo": s(c), "descripcion": (d or "").strip()}
        for c, d in col(None, 2, 504, [ci("AX"), ci("AY")])
    ]
    # En el Excel la lista desplegable solo alcanza BB2:BB1109; se marcan las demás.
    arl = []
    for r in range(2, 1672):
        d, c = v.cell(r, ci("BA")).value, v.cell(r, ci("BB")).value
        if c is None:
            continue
        arl.append({"codigo": s(c), "descripcion": " ".join((d or "").split()), "enListaExcel": r <= 1109})

    motivos_dep = []
    for r in range(128, 133):
        motivos_dep.append({
            "texto": ded.cell(r, ci("C")).value,
            "numero": ded.cell(r, ci("P")).value,
            "anexo": ded.cell(r, ci("Q")).value,
        })
    parentescos = [ded.cell(r, ci("J")).value for r in range(128, 135) if ded.cell(r, ci("J")).value]

    data = {
        "_fuente": XLSX.name,
        "tiposDocumento": tipos_doc,
        "paises": paises,
        "departamentos": departamentos,
        "municipios": municipios,
        "bancos": bancos,
        "tiposCuenta": tipos_cuenta,
        "regimenIva": regimen,
        "areas": areas,
        "rit": rit,
        "ciiu": ciiu,
        "arl": arl,
        "motivosDependientes": motivos_dep,
        "parentescos": parentescos,
        "juramento": v["D69"].value,
        "declaracionDeducciones": ded["B60"].value,
        "autorizacionDeducciones": ded["B70"].value,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    body = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    OUT.write_text(
        "// Generado por tools/extract_catalogos.py a partir del Excel original. No editar a mano.\n"
        f"window.CATALOGOS = {body};\n",
        encoding="utf-8",
    )
    print(f"OK -> {OUT} ({OUT.stat().st_size // 1024} KB)")
    for k, val in data.items():
        if isinstance(val, list):
            print(f"  {k}: {len(val)}")


if __name__ == "__main__":
    main()
